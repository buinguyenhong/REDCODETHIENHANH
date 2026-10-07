import io
from datetime import datetime, timezone, timedelta
from typing import Optional, Dict, Any, Tuple
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select, func, desc
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from zoneinfo import ZoneInfo
from app.database import get_db
from app.models import Alarm, AlarmType, Department, Station, SystemEvent, User, AlarmStationState
from app.api.deps import get_current_active_admin

router = APIRouter(prefix="/reports", tags=["Reports"])

def parse_date_range(from_date_str: Optional[str], to_date_str: Optional[str]) -> Tuple[Optional[datetime], Optional[datetime]]:
    """
    Half-open [start, end): date-only end is next hospital midnight.
    Naive datetimes are hospital local time; offset datetimes are normalized UTC.
    """
    start_dt = None
    end_dt = None
    if from_date_str and from_date_str.strip():
        clean_str = from_date_str.strip()
        try:
            dt = datetime.fromisoformat(clean_str)
        except ValueError:
            raise HTTPException(status_code=422, detail='Ngày không hợp lệ')
        if len(clean_str) <= 10:
            dt = dt.replace(hour=0, minute=0, second=0, microsecond=0, tzinfo=ZoneInfo('Asia/Ho_Chi_Minh')).astimezone(timezone.utc)
        elif dt.tzinfo is None:
            dt = dt.replace(tzinfo=ZoneInfo('Asia/Ho_Chi_Minh'))
        dt = dt.astimezone(timezone.utc)
        start_dt = dt

    if to_date_str and to_date_str.strip():
        clean_str = to_date_str.strip()
        try:
            dt = datetime.fromisoformat(clean_str)
        except ValueError:
            raise HTTPException(status_code=422, detail='Ngày không hợp lệ')
        if len(clean_str) <= 10:
            dt = (dt.replace(hour=0, minute=0, second=0, microsecond=0, tzinfo=ZoneInfo('Asia/Ho_Chi_Minh')) + timedelta(days=1)).astimezone(timezone.utc)
        elif dt.tzinfo is None:
            dt = dt.replace(tzinfo=ZoneInfo('Asia/Ho_Chi_Minh'))
        dt = dt.astimezone(timezone.utc)
        end_dt = dt

    if start_dt and end_dt and start_dt >= end_dt:
        raise HTTPException(status_code=422, detail='Khoảng ngày không hợp lệ')
    return start_dt, end_dt

@router.get("/summary")
async def get_summary_report(
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_admin)
):
    start_dt, end_dt = parse_date_range(from_date, to_date)

    stmt = select(Alarm).options(
        selectinload(Alarm.alarm_type),
        selectinload(Alarm.source_department)
    )
    if start_dt:
        stmt = stmt.where(Alarm.created_at >= start_dt)
    if end_dt:
        stmt = stmt.where(Alarm.created_at < end_dt)

    alarms = (await db.execute(stmt)).scalars().all()

    total_alarms = len(alarms)
    by_type: Dict[str, int] = {}
    by_dept: Dict[str, int] = {}
    by_status: Dict[str, int] = {}
    by_day, by_month = {}, {}

    for a in alarms:
        local = a.created_at.replace(tzinfo=timezone.utc).astimezone(ZoneInfo('Asia/Ho_Chi_Minh')) if a.created_at.tzinfo is None else a.created_at.astimezone(ZoneInfo('Asia/Ho_Chi_Minh'))
        day, month = local.strftime('%Y-%m-%d'), local.strftime('%Y-%m')
        by_day[day] = by_day.get(day, 0) + 1
        by_month[month] = by_month.get(month, 0) + 1
        t_name = a.alarm_type.name if a.alarm_type else "Không xác định"
        d_name = a.source_department.name if a.source_department else "Không xác định"
        by_type[t_name] = by_type.get(t_name, 0) + 1
        by_dept[d_name] = by_dept.get(d_name, 0) + 1
        by_status[a.status] = by_status.get(a.status, 0) + 1

    # Filter offline events with the exact same date range (Section 28)
    offline_stmt = select(func.count(SystemEvent.id)).where(SystemEvent.event_type == "DEVICE_DISCONNECTED")
    if start_dt:
        offline_stmt = offline_stmt.where(SystemEvent.created_at >= start_dt)
    if end_dt:
        offline_stmt = offline_stmt.where(SystemEvent.created_at < end_dt)

    offline_count = (await db.execute(offline_stmt)).scalar() or 0

    return {
        "total_alarms": total_alarms,
        "by_day": by_day,
        "by_month": by_month,
        "by_type": by_type,
        "by_department": by_dept,
        "by_status": by_status,
        "device_offline_events_count": offline_count
    }

@router.get("/export-xlsx")
async def export_alarms_xlsx(
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_admin)
):
    start_dt, end_dt = parse_date_range(from_date, to_date)

    stmt = (
        select(Alarm)
        .options(
            selectinload(Alarm.alarm_type),
            selectinload(Alarm.source_department),
            selectinload(Alarm.created_by_user),
            selectinload(Alarm.station_states).selectinload(AlarmStationState.station)
        )
        .order_by(Alarm.created_at.asc())
    )
    if start_dt:
        stmt = stmt.where(Alarm.created_at >= start_dt)
    if end_dt:
        stmt = stmt.where(Alarm.created_at < end_dt)

    alarms = (await db.execute(stmt)).scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "BaoCao_Redcode"
    details = wb.create_sheet('TrangThai_Tram')
    details.append(['Alarm ID', 'Trạm', 'State', 'Received', 'Displayed', 'Audio started', 'Audio completed', 'Dismissed', 'Failed', 'Error'])
    def safe(value):
        return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value
    def formatted(value):
        if not value:
            return ''
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        return value.astimezone(ZoneInfo('Asia/Ho_Chi_Minh')).isoformat()

    # Styling definitions
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "STT", "Thời gian", "Mã báo động", "Tên báo động",
        "Khoa/Phòng", "Vị trí", "Ghi chú", "Người phát", "Trạng thái", "Thứ tự Server", "Chi tiết trạm nhận (States)"
    ]
    ws.append(headers)

    for col_idx, col_name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    ws.row_dimensions[1].height = 26

    for idx, a in enumerate(alarms, 1):
        created_str = formatted(a.created_at)
        t_code = a.alarm_type.code if a.alarm_type else ""
        t_name = a.alarm_type.name if a.alarm_type else ""
        dept_name = a.source_department.name if a.source_department else ""
        user_name = a.created_by_user.display_name if a.created_by_user else ""

        # Summarize station states
        station_summaries = []
        if a.station_states:
            for st in a.station_states:
                st_code = st.station.station_code if st.station else f"ID-{st.station_id}"
                details.append([safe(v) for v in [a.id, st_code, st.state, formatted(st.received_at), formatted(st.displayed_at), formatted(st.audio_started_at), formatted(st.audio_completed_at), formatted(st.dismissed_at), formatted(st.failed_at), st.error_message]])
                st_text = f"{st_code}: {st.state}"
                if st.dismissed_at:
                    st_text += f" (Dismiss: {st.dismissed_at.strftime('%H:%M:%S')})"
                station_summaries.append(st_text)
        stations_str = "; ".join(station_summaries) if station_summaries else "Chưa có trạm phản hồi"

        row_data = [
            idx,
            created_str,
            t_code,
            t_name,
            dept_name,
            a.source_location,
            a.note,
            user_name,
            a.status,
            a.server_sequence,
            stations_str
        ]
        ws.append([safe(v) for v in row_data])

        row_num = idx + 1
        ws.row_dimensions[row_num].height = 20
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 9, 10]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Column widths
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 26
    ws.column_dimensions['E'].width = 24
    ws.column_dimensions['F'].width = 26
    ws.column_dimensions['G'].width = 32
    ws.column_dimensions['H'].width = 20
    ws.column_dimensions['I'].width = 14
    ws.column_dimensions['J'].width = 14
    ws.column_dimensions['K'].width = 40

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"BaoCao_Redcode_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
