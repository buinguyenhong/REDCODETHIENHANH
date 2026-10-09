import io
import json
import os
import uuid
from datetime import datetime, timezone, timedelta
from unittest.mock import AsyncMock, patch

import pytest
import pytest_asyncio
import websockets
from httpx import AsyncClient, ASGITransport
from sqlalchemy import select, func
from openpyxl import load_workbook

from app.main import app
from app.database import Base, engine, AsyncSessionLocal
from app.models import Alarm, AlarmEvent, AlarmStationState, NotificationOutbox, SystemEvent
from app.seed import seed_database
from app.core.station_lifecycle import acknowledge_station_event
from app.core.alarm_lifecycle import expire_alarms
from app.core.n8n_outbox import outbox_worker
from app.config import settings
from app.api.reports import parse_date_range
from tests.live_server import running_server


@pytest_asyncio.fixture(scope='module', autouse=True)
async def setup():
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    await seed_database()


async def admin_headers(client):
    response = await client.post('/api/auth/login', json={'username': 'admin', 'password': 'admin123456'})
    return {'Authorization': f"Bearer {response.json()['access_token']}"}


async def receiver_alarm(client, headers):
    station = (await client.post('/api/stations/register', headers=headers, json={
        'station_code': 'LIFE-' + uuid.uuid4().hex[:10], 'name': 'Lifecycle', 'receiver_group_ids': [1]
    })).json()
    alarm = (await client.post('/api/alarms', headers=headers, json={'alarm_type_id': 2})).json()
    return station, alarm


@pytest.mark.asyncio
async def test_station_full_lifecycle_idempotent_and_no_regression():
    async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
        headers = await admin_headers(client)
        station, alarm = await receiver_alarm(client, headers)
        steps = [('RECEIVED', 'DELIVERED', 'received_at'), ('DISPLAYED', 'DISPLAYED', 'displayed_at'),
                 ('AUDIO_STARTED', 'AUDIO_STARTED', 'audio_started_at'), ('AUDIO_COMPLETED', 'AUDIO_COMPLETED', 'audio_completed_at')]
        async with AsyncSessionLocal() as session:
            assert not await acknowledge_station_event(session, alarm['id'], station['id'], 'AUDIO_STARTED', {})
            assert not await acknowledge_station_event(session, alarm['id'], station['id'], 'DISPLAYED', {})
        for event, expected, timestamp in steps:
            async with AsyncSessionLocal() as session:
                assert await acknowledge_station_event(session, alarm['id'], station['id'], event, {})
                target = (await session.execute(select(AlarmStationState).where(
                    AlarmStationState.alarm_id == alarm['id'], AlarmStationState.station_id == station['id']))).scalar_one()
                first_time = getattr(target, timestamp)
                assert target.state == expected and first_time is not None
                assert not await acknowledge_station_event(session, alarm['id'], station['id'], event, {})
                await session.refresh(target)
                assert getattr(target, timestamp) == first_time
                if event != 'RECEIVED':
                    assert not await acknowledge_station_event(session, alarm['id'], station['id'], 'RECEIVED', {})
        dismiss = {'alarm_id': alarm['id'], 'station_code': station['station_code'], 'device_token': station['raw_device_token']}
        assert (await client.post('/api/stations/dismiss', json=dismiss)).status_code == 200
        assert (await client.post('/api/stations/dismiss', json=dismiss)).status_code == 200
        async with AsyncSessionLocal() as session:
            assert not await acknowledge_station_event(session, alarm['id'], station['id'], 'AUDIO_FAILED', {'error': 'late'})
            assert (await session.get(Alarm, alarm['id'])).status == 'ACTIVE'
            events = (await session.execute(select(AlarmEvent.event_type).where(
                AlarmEvent.alarm_id == alarm['id'], AlarmEvent.station_id == station['id']))).scalars().all()
            assert events == [step[0] for step in steps] + ['DISMISSED']


@pytest.mark.asyncio
async def test_audio_failure_preserves_error_and_requires_successful_retry():
    async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
        station, alarm = await receiver_alarm(client, await admin_headers(client))
        async with AsyncSessionLocal() as session:
            for event in ['RECEIVED', 'DISPLAYED', 'AUDIO_FAILED']:
                assert await acknowledge_station_event(session, alarm['id'], station['id'], event, {'error': 'autoplay blocked'})
            target = (await session.execute(select(AlarmStationState).where(
                AlarmStationState.alarm_id == alarm['id'], AlarmStationState.station_id == station['id']))).scalar_one()
            assert target.state == 'FAILED' and target.error_message == 'autoplay blocked'
            assert target.audio_started_at is None and target.audio_completed_at is None
            assert not await acknowledge_station_event(session, alarm['id'], station['id'], 'AUDIO_COMPLETED', {})
            assert not await acknowledge_station_event(session, alarm['id'], station['id'], 'AUDIO_FAILED', {'error': 'duplicate'})
            assert await acknowledge_station_event(session, alarm['id'], station['id'], 'AUDIO_STARTED', {})
            assert await acknowledge_station_event(session, alarm['id'], station['id'], 'AUDIO_COMPLETED', {})
            await session.refresh(target)
            assert target.failed_at and target.error_message == 'autoplay blocked'


@pytest.mark.asyncio
async def test_terminal_global_states_and_target_broadcast_once():
    async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
        headers = await admin_headers(client)
        station, alarm = await receiver_alarm(client, headers)
        with patch('app.core.websocket_manager.manager.send_to_station', new_callable=AsyncMock) as send, patch('app.core.websocket_manager.manager.broadcast_to_dashboards', new_callable=AsyncMock) as dashboard:
            assert (await client.post(f"/api/alarms/{alarm['id']}/cancel", headers=headers)).json()['status'] == 'CANCELLED'
            assert (await client.post(f"/api/alarms/{alarm['id']}/cancel", headers=headers)).status_code == 200
            assert sum(call.args[0] == station['station_code'] and call.args[1]['type'] == 'ALARM_CANCELLED' for call in send.call_args_list) == 1
            assert any(call.args[0]['type'] == 'ALARM_CANCELLED' for call in dashboard.call_args_list)
        async with AsyncSessionLocal() as session:
            cancelled = await session.get(Alarm, alarm['id'])
            cancelled.expires_at = datetime.now(timezone.utc) - timedelta(seconds=1)
            await session.commit()
            assert alarm['id'] not in await expire_alarms(session)
            await session.refresh(cancelled)
            assert cancelled.status == 'CANCELLED'
        station, alarm = await receiver_alarm(client, headers)
        async with AsyncSessionLocal() as session:
            stored = await session.get(Alarm, alarm['id'])
            stored.expires_at = datetime.now(timezone.utc) - timedelta(seconds=1)
            await session.commit()
        with patch('app.core.websocket_manager.manager.send_to_station', new_callable=AsyncMock) as send, patch('app.core.websocket_manager.manager.broadcast_to_dashboards', new_callable=AsyncMock) as dashboard:
            assert (await client.get(f"/api/alarms/{alarm['id']}", headers=headers)).json()['status'] == 'EXPIRED'
            assert (await client.post(f"/api/alarms/{alarm['id']}/cancel", headers=headers)).status_code == 409
            assert sum(call.args[0] == station['station_code'] and call.args[1]['type'] == 'ALARM_EXPIRED' for call in send.call_args_list) == 1
            assert any(call.args[0]['type'] == 'ALARM_EXPIRED' for call in dashboard.call_args_list)
        async with AsyncSessionLocal() as session:
            assert not await acknowledge_station_event(session, alarm['id'], station['id'], 'RECEIVED', {})
            assert (await session.execute(select(func.count(AlarmEvent.id)).where(AlarmEvent.alarm_id == alarm['id'], AlarmEvent.event_type == 'EXPIRED'))).scalar() == 1
            assert (await session.execute(select(func.count(NotificationOutbox.id)).where(NotificationOutbox.alarm_id == alarm['id']))).scalar() == 1


@pytest.mark.asyncio
async def test_permission_single_source_updates_missing_denied_and_spoof():
    async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
        headers = await admin_headers(client)
        operator = (await client.post('/api/auth/login', json={'username': 'operator_cc', 'password': 'pass123456'})).json()['access_token']
        op_headers = {'Authorization': f'Bearer {operator}'}
        created = (await client.post('/api/alarm-types', headers=headers, json={
            'code': 'PERM-' + uuid.uuid4().hex[:8], 'name': 'Permission test', 'allowed_department_ids': [2]
        })).json()
        path = f"/api/alarm-types/{created['id']}"
        payload = {'alarm_type_id': created['id'], 'source_department_id': 9}
        triggered = await client.post('/api/alarms', headers=op_headers, json=payload)
        assert triggered.status_code == 201 and triggered.json()['source_department_id'] == 2
        assert (await client.put(path, headers=headers, json={'allowed_department_ids': []})).json()['allowed_department_ids'] == []
        assert (await client.post('/api/alarms', headers=op_headers, json=payload)).status_code == 403
        from app.models import DepartmentAlarmPermission
        async with AsyncSessionLocal() as session:
            await session.execute(DepartmentAlarmPermission.__table__.delete().where(DepartmentAlarmPermission.alarm_type_id == created['id']))
            await session.commit()
        assert (await client.post('/api/alarms', headers=op_headers, json=payload)).status_code == 403
        assert (await client.put(path, headers=headers, json={'allowed_department_ids': [2]})).status_code == 200
        assert (await client.post('/api/alarms', headers=op_headers, json=payload)).status_code == 201
        assert (await client.put(path, headers=headers, json={'allowed_department_ids': [999999]})).status_code == 422


@pytest.mark.parametrize('start,end,next_day', [
    ('2026-10-06', '2026-10-06', '2026-10-07'),
    ('2026-10-01', '2026-10-06', '2026-10-07'),
    ('2026-10-31', '2026-10-31', '2026-11-01'),
    ('2026-12-31', '2026-12-31', '2027-01-01'),
])
@pytest.mark.asyncio
async def test_report_half_open_boundaries_summary_offline_and_export(start, end, next_day):
    start_dt, end_dt = parse_date_range(start, end)
    assert end_dt == datetime.fromisoformat(next_day).replace(tzinfo=timezone(timedelta(hours=7))).astimezone(timezone.utc)
    times = [start_dt - timedelta(microseconds=1), start_dt, end_dt - timedelta(microseconds=1), end_dt]
    async with AsyncSessionLocal() as session:
        alarms = [Alarm(alarm_type_id=2, status='EXPIRED', source_location='boundary', note='', server_sequence=0, created_at=t) for t in times]
        session.add_all(alarms)
        session.add_all([SystemEvent(event_type='DEVICE_DISCONNECTED', severity='WARNING', message='boundary', created_at=t) for t in times])
        await session.commit()
        expected_ids = {alarms[1].id, alarms[2].id}
        excluded_ids = {alarms[0].id, alarms[3].id}
    async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
        headers = await admin_headers(client)
        params = {'from_date': start, 'to_date': end}
        summary = (await client.get('/api/reports/summary', headers=headers, params=params)).json()
        assert summary['total_alarms'] >= 2 and summary['device_offline_events_count'] >= 2
        response = await client.get('/api/reports/export-xlsx', headers=headers, params=params)
        workbook = load_workbook(io.BytesIO(response.content))
        # Detail sheet has no station rows for synthetic alarms; compare main
        # rows against SQL selection for exact inclusion/exclusion.
        async with AsyncSessionLocal() as session:
            ids = set((await session.execute(select(Alarm.id).where(Alarm.created_at >= start_dt, Alarm.created_at < end_dt))).scalars())
            offline = (await session.execute(select(func.count(SystemEvent.id)).where(SystemEvent.event_type == 'DEVICE_DISCONNECTED', SystemEvent.created_at >= start_dt, SystemEvent.created_at < end_dt))).scalar()
        assert expected_ids <= ids and not excluded_ids & ids
        assert summary['total_alarms'] == len(ids)
        assert summary['device_offline_events_count'] == offline
        assert workbook.active.max_row - 1 == len(ids)


@pytest.mark.asyncio
async def test_station_ws_rejects_missing_wrong_credentials():
    """A rejected station credential must arrive as close code 1008, not a 403.

    Closing before accept() makes the server answer the upgrade with HTTP 403,
    which browsers report as close code 1006 with no reason — indistinguishable
    from a dropped network, so the kiosk would retry forever instead of asking
    to be re-confirmed. The gateway therefore accepts and closes with 1008.

    Runs against a real uvicorn: the 403-vs-1008 choice is made by the ASGI
    server itself, so TestClient cannot stand in for it here.
    """
    async with running_server(os.getenv('DATABASE_URL', 'sqlite+aiosqlite:///./redcode.db')) as (_, ws_url):
        for query in ['type=station', 'type=station&station_code=ST-CC-01&token=wrong']:
            async with websockets.connect(f'{ws_url}?{query}') as ws:
                with pytest.raises(websockets.exceptions.ConnectionClosed) as exc:
                    await ws.recv()
            assert exc.value.rcvd.code == 1008
            assert exc.value.rcvd.reason, 'client must receive a readable reason'


@pytest.mark.asyncio
async def test_outbox_claim_crash_recovery_and_duplicate_claim():
    async with AsyncSessionLocal() as session:
        item = NotificationOutbox(event_type='ALARM_CREATED', payload={}, status='PENDING', attempt_count=0)
        session.add(item)
        await session.commit()
        item_id = item.id
    class Crash(BaseException):
        pass
    with patch('httpx.AsyncClient.post', new_callable=AsyncMock, side_effect=Crash()):
        with pytest.raises(Crash):
            await outbox_worker._send_item(item_id)
    async with AsyncSessionLocal() as session:
        item = await session.get(NotificationOutbox, item_id)
        assert item.status == 'PROCESSING' and item.locked_at and item.attempt_count == 1
        item.locked_at = datetime.now(timezone.utc) - timedelta(minutes=6)
        await session.commit()
    # Reclaim runs even if this job is beyond the first due batch.
    with patch.object(outbox_worker, '_send_item', new_callable=AsyncMock):
        await outbox_worker.process_outbox()
    with patch('httpx.AsyncClient.post', new_callable=AsyncMock, side_effect=RuntimeError('n8n down')) as send:
        await outbox_worker._send_item(item_id)
        await outbox_worker._send_item(item_id)
        assert send.await_count == 1
    async with AsyncSessionLocal() as session:
        item = await session.get(NotificationOutbox, item_id)
        assert item.status == 'FAILED' and item.attempt_count == 2 and item.locked_at is None
        assert item.next_attempt_at is not None and 'n8n down' in item.last_error


@pytest.mark.asyncio
async def test_idempotency_default_location_and_source_fingerprint():
    async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
        headers = await admin_headers(client)
        request = {'alarm_type_id': 2, 'idempotency_key': str(uuid.uuid4())}
        first = await client.post('/api/alarms', headers=headers, json=request)
        retry = await client.post('/api/alarms', headers=headers, json=request)
        assert first.status_code == retry.status_code == 201
        assert first.json()['id'] == retry.json()['id']
        changed = await client.post('/api/alarms', headers=headers, json={**request, 'source_department_id': 9})
        assert changed.status_code == 409


@pytest.mark.asyncio
async def test_production_startup_never_calls_create_all():
    original = settings.ENVIRONMENT
    try:
        settings.ENVIRONMENT = 'production'
        with patch.object(Base.metadata, 'create_all', side_effect=AssertionError('Production schema mutation')):
            async with app.router.lifespan_context(app):
                async with AsyncSessionLocal() as session:
                    assert (await session.execute(select(func.count(Alarm.id)))).scalar() >= 0
    finally:
        settings.ENVIRONMENT = original
